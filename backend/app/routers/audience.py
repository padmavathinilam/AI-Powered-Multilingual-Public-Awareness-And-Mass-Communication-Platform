
import csv
import io
import re
from math import ceil

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.audience_member import AudienceMember
from app.schemas.audience import (
    AudienceCreate,
    AudienceUpdate,
    LANGUAGES,
    CHANNELS,
    CONSENT_STATUSES,
)

router = APIRouter(prefix="/api/audience", tags=["Audience"])

MAX_UPLOAD_SIZE = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


def member_to_dict(member: AudienceMember):
    return {
        "id": member.id,
        "name": member.name,
        "email": member.email,
        "phone": member.phone,
        "age": member.age,
        "gender": member.gender,
        "city": member.city,
        "state": member.state,
        "preferred_language": member.preferred_language,
        "occupation": member.occupation,
        "organization_department": member.organization_department,
        "preferred_channel": member.preferred_channel,
        "consent_status": member.consent_status,
        "created_at": (
            member.created_at.isoformat()
            if member.created_at
            else None
        ),
    }


def validate_choices(payload):
    if payload.preferred_language not in LANGUAGES:
        raise HTTPException(
            status_code=422,
            detail=f"Unsupported language. Choose from: {LANGUAGES}",
        )

    if payload.preferred_channel not in CHANNELS:
        raise HTTPException(
            status_code=422,
            detail=f"Unsupported channel. Choose from: {CHANNELS}",
        )

    if payload.consent_status not in CONSENT_STATUSES:
        raise HTTPException(
            status_code=422,
            detail=f"Invalid consent status. Choose from: {CONSENT_STATUSES}",
        )


def find_duplicate(db: Session, email=None, phone=None, exclude_id=None):
    conditions = []

    if email:
        conditions.append(
            func.lower(AudienceMember.email) == email.strip().lower()
        )

    if phone:
        conditions.append(
            AudienceMember.phone == phone.strip()
        )

    if not conditions:
        return None

    query = db.query(AudienceMember).filter(or_(*conditions))

    if exclude_id is not None:
        query = query.filter(AudienceMember.id != exclude_id)

    return query.first()


def grouped_counts(db: Session, field):
    rows = (
        db.query(field, func.count(AudienceMember.id))
        .group_by(field)
        .all()
    )

    return {
        value if value else "Not specified": count
        for value, count in rows
    }


def normalize_header(value):
    value = str(value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")

    aliases = {
        "full_name": "name",
        "mobile": "phone",
        "mobile_number": "phone",
        "phone_number": "phone",
        "email_address": "email",
        "preferred_lang": "preferred_language",
        "language": "preferred_language",
        "department": "organization_department",
        "organization": "organization_department",
        "channel": "preferred_channel",
        "consent": "consent_status",
    }

    return aliases.get(value, value)


def clean_import_row(row):
    allowed = {
        "name",
        "email",
        "phone",
        "age",
        "gender",
        "city",
        "state",
        "preferred_language",
        "occupation",
        "organization_department",
        "preferred_channel",
        "consent_status",
    }

    cleaned = {}

    for key, value in row.items():
        key = normalize_header(key)

        if key not in allowed:
            continue

        if value is None:
            cleaned[key] = None
        elif isinstance(value, str):
            cleaned[key] = value.strip() or None
        else:
            cleaned[key] = value

    return cleaned


def read_import_rows(filename, content):
    extension = filename.lower().rsplit(".", 1)[-1]

    if extension == "csv":
        try:
            text_content = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=400,
                detail="CSV must use UTF-8 encoding.",
            )

        reader = csv.DictReader(io.StringIO(text_content))

        if not reader.fieldnames:
            raise HTTPException(
                status_code=400,
                detail="The CSV file must contain a header row.",
            )

        headers = [normalize_header(h) for h in reader.fieldnames]
        required = {
            "name",
            "email",
            "phone",
            "age",
            "gender",
            "city",
            "state",
            "preferred_language",
            "occupation",
            "preferred_channel",
            "consent_status",
        }

        missing = required - set(headers)

        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing columns: {', '.join(sorted(missing))}",
            )

        return [
            clean_import_row(row)
            for row in reader
        ]

    if extension == "xlsx":
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )

            try:
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    raise HTTPException(
                        status_code=400,
                        detail="The Excel file must contain a header row.",
                    )

                headers = [normalize_header(h) for h in headers]

                required = {
                    "name",
                    "email",
                    "phone",
                    "age",
                    "gender",
                    "city",
                    "state",
                    "preferred_language",
                    "occupation",
                    "preferred_channel",
                    "consent_status",
                }

                missing = required - set(headers)

                if missing:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Missing columns: {', '.join(sorted(missing))}",
                    )

                return [
                    clean_import_row(dict(zip(headers, values)))
                    for values in rows
                ]
            finally:
                workbook.close()

        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read Excel file: {str(exc)}",
            )

    raise HTTPException(
        status_code=400,
        detail="Upload a .csv or .xlsx file.",
    )


# GET audience members: search, filters and pagination
@router.get("/")
def get_audience_members(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    q: str | None = Query(None, max_length=150),
    city: str | None = None,
    state: str | None = None,
    preferred_language: str | None = None,
    preferred_channel: str | None = None,
    consent_status: str | None = None,
    occupation: str | None = None,
    db: Session = Depends(get_db),
):
    query = db.query(AudienceMember)

    if q and q.strip():
        search = f"%{q.strip()}%"
        query = query.filter(
            or_(
                AudienceMember.name.ilike(search),
                AudienceMember.email.ilike(search),
                AudienceMember.phone.ilike(search),
                AudienceMember.city.ilike(search),
            )
        )

    if city and city.strip():
        query = query.filter(
            func.lower(AudienceMember.city) == city.strip().lower()
        )

    if state and state.strip():
        query = query.filter(
            func.lower(AudienceMember.state) == state.strip().lower()
        )

    if preferred_language and preferred_language.strip():
        query = query.filter(
            func.lower(AudienceMember.preferred_language)
            == preferred_language.strip().lower()
        )

    if preferred_channel and preferred_channel.strip():
        query = query.filter(
            func.lower(AudienceMember.preferred_channel)
            == preferred_channel.strip().lower()
        )

    if occupation and occupation.strip():
        query = query.filter(
            func.lower(AudienceMember.occupation)
            == occupation.strip().lower()
        )

    if consent_status and consent_status.strip():
        value = consent_status.strip().lower()

        valid_status = next(
            (
                item
                for item in CONSENT_STATUSES
                if item.lower() == value
            ),
            None,
        )

        if valid_status is None:
            raise HTTPException(
                status_code=422,
                detail=f"Invalid consent status. Choose from: {CONSENT_STATUSES}",
            )

        query = query.filter(
            func.lower(AudienceMember.consent_status)
            == valid_status.lower()
        )

    total = query.count()

    members = (
        query.order_by(AudienceMember.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    return {
        "items": [member_to_dict(member) for member in members],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": ceil(total / page_size) if total else 0,
    }


# Audience statistics
@router.get("/stats")
def get_audience_stats(db: Session = Depends(get_db)):
    return {
        "total": db.query(AudienceMember).count(),
        "by_language": grouped_counts(
            db, AudienceMember.preferred_language
        ),
        "by_channel": grouped_counts(
            db, AudienceMember.preferred_channel
        ),
        "by_consent_status": grouped_counts(
            db, AudienceMember.consent_status
        ),
        "by_city": grouped_counts(db, AudienceMember.city),
    }


# Bulk import from CSV or Excel
@router.post("/bulk-upload")
async def bulk_upload_audience(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    filename = file.filename or ""

    if "." not in filename:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file must be CSV or XLSX.",
        )

    extension = filename.lower().rsplit(".", 1)[-1]

    if extension not in {"csv", "xlsx"}:
        raise HTTPException(
            status_code=400,
            detail="Upload a .csv or .xlsx file.",
        )

    content = await file.read(MAX_UPLOAD_SIZE + 1)
    await file.close()

    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File size must not exceed 10 MB.",
        )

    rows = read_import_rows(filename, content)

    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=400,
            detail=f"Import is limited to {MAX_IMPORT_ROWS} rows per file.",
        )

    imported = 0
    skipped = 0
    errors = []
    seen_emails = set()
    seen_phones = set()
    pending_members = []

    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None for value in row.values()):
            skipped += 1
            continue

        try:
            payload = AudienceCreate.model_validate(row)
            validate_choices(payload)

            email = str(payload.email).strip().lower()
            phone = str(payload.phone).strip()

            if email in seen_emails:
                skipped += 1
                errors.append({
                    "row": row_number,
                    "reason": "Duplicate email in uploaded file.",
                })
                continue

            if phone in seen_phones:
                skipped += 1
                errors.append({
                    "row": row_number,
                    "reason": "Duplicate phone number in uploaded file.",
                })
                continue

            if find_duplicate(db, email=email, phone=phone):
                skipped += 1
                errors.append({
                    "row": row_number,
                    "reason": "Email or phone already exists in the database.",
                })
                continue

            pending_members.append(
                AudienceMember(**payload.model_dump())
            )
            seen_emails.add(email)
            seen_phones.add(phone)
            imported += 1

        except Exception as exc:
            skipped += 1
            reason = getattr(exc, "errors", None)
            errors.append({
                "row": row_number,
                "reason": str(reason() if callable(reason) else exc)[:500],
            })

    try:
        db.add_all(pending_members)
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Bulk import failed. No imported records were saved.",
        )

    return {
        "message": "Bulk import completed.",
        "imported": imported,
        "skipped": skipped,
        "errors": errors[:100],
        "errors_truncated": len(errors) > 100,
    }


# Create audience member
@router.post("/", status_code=status.HTTP_201_CREATED)
def create_audience_member(
    payload: AudienceCreate,
    db: Session = Depends(get_db),
):
    validate_choices(payload)

    if find_duplicate(db, email=payload.email, phone=payload.phone):
        raise HTTPException(
            status_code=409,
            detail="A member with this email or phone number already exists.",
        )

    member = AudienceMember(**payload.model_dump())
    db.add(member)

    try:
        db.commit()
        db.refresh(member)
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="A member with this email or phone number already exists.",
        )

    return member_to_dict(member)


# Get one audience member
@router.get("/{member_id}")
def get_audience_member(
    member_id: int,
    db: Session = Depends(get_db),
):
    member = db.query(AudienceMember).filter(
        AudienceMember.id == member_id
    ).first()

    if not member:
        raise HTTPException(
            status_code=404,
            detail="Audience member not found.",
        )

    return member_to_dict(member)


# Update an audience member
@router.put("/{member_id}")
def update_audience_member(
    member_id: int,
    payload: AudienceUpdate,
    db: Session = Depends(get_db),
):
    member = db.query(AudienceMember).filter(
        AudienceMember.id == member_id
    ).first()

    if not member:
        raise HTTPException(
            status_code=404,
            detail="Audience member not found.",
        )

    validate_choices(payload)

    if find_duplicate(
        db,
        email=payload.email,
        phone=payload.phone,
        exclude_id=member_id,
    ):
        raise HTTPException(
            status_code=409,
            detail="Another member already uses this email or phone number.",
        )

    for field, value in payload.model_dump().items():
        setattr(member, field, value)

    try:
        db.commit()
        db.refresh(member)
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Another member already uses this email or phone number.",
        )

    return member_to_dict(member)


# Delete an audience member
@router.delete("/{member_id}")
def delete_audience_member(
    member_id: int,
    db: Session = Depends(get_db),
):
    member = db.query(AudienceMember).filter(
        AudienceMember.id == member_id
    ).first()

    if not member:
        raise HTTPException(
            status_code=404,
            detail="Audience member not found.",
        )

    db.delete(member)
    db.commit()

    return {
        "message": "Audience member deleted successfully.",
        "id": member_id,
    }
